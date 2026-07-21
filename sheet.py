"""Read-only access to Nurin's outreach sheet — the dedupe source.

Fail-closed (spec §9-F3): any doubt about the tab/column/credentials BLOCKS
discovery. Surfacing an already-contacted creator is the one thing this tool
must never do, so there is no 'proceed anyway'.

The outreach file can be either a native Google Sheet or an Office (.xlsx)
file living in Drive — Nurin's is the latter, and the Sheets API refuses to
read Office files, so we ask Drive for the file's type and pick the right
access path. Both the Sheets API and the Drive API must be enabled on the
service account's Google Cloud project (README §4).
"""
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

import requests

APP_DIR = Path(__file__).parent
_SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly",
           "https://www.googleapis.com/auth/drive.readonly"]

_GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The real header may sit below banner rows (Nurin's tab has a
# "Working Year / Working Month" row above it). Scan this many rows for an
# exact header match; anything above the header is ignored, no header found
# still blocks.
HEADER_SCAN_ROWS = 10


class SheetError(RuntimeError):
    pass


def parse_contacted(values, expected_header, where):
    want = str(expected_header).strip().lower()
    header_idx = None
    for i, row in enumerate(values[:HEADER_SCAN_ROWS]):
        cell = "" if not row or row[0] is None else str(row[0]).strip()
        if cell.lower() == want:
            header_idx = i
            break
    if header_idx is None:
        seen = [str(r[0]).strip() for r in values[:HEADER_SCAN_ROWS]
                if r and r[0] is not None and str(r[0]).strip()]
        raise SheetError(
            f"expected header '{expected_header}' in the first {HEADER_SCAN_ROWS} rows of {where}, "
            f"found {seen[:5]!r}. Wrong tab/column in config.json?")
    # Empty xlsx cells arrive as None — they must vanish, never become "None".
    handles = [str(v[0]).strip() for v in values[header_idx + 1:]
               if v and v[0] is not None and str(v[0]).strip()]
    if not handles:
        raise SheetError(
            f"0 contacted handles under {where} — almost certainly the wrong column; "
            "refusing to run discovery without a dedupe source.")
    return handles


def _get(url, token, timeout=30):
    try:
        return requests.get(url, headers={"Authorization": f"Bearer {token}"}, timeout=timeout)
    except requests.RequestException as e:
        raise SheetError(f"could not reach Google: {e}") from e


def _json(resp, what):
    try:
        return resp.json()
    except ValueError as e:
        raise SheetError(f"{what} returned a non-JSON response: {e}") from e


def _file_mime(sheet_id, token):
    r = _get(f"https://www.googleapis.com/drive/v3/files/{sheet_id}"
             "?fields=mimeType&supportsAllDrives=true", token)
    if r.status_code != 200:
        raise SheetError(
            f"Drive API HTTP {r.status_code} looking up the outreach file: {r.text[:200]} — "
            "is the Drive API enabled and the file shared with the service account?")
    return _json(r, "Drive API").get("mimeType", "")


def _values_from_sheets_api(sheet_cfg, token):
    tab, col = sheet_cfg["outreach_tab"], sheet_cfg["handle_column"]
    range_ref = f"'{tab}'!{col}:{col}"
    r = _get(f"https://sheets.googleapis.com/v4/spreadsheets/{sheet_cfg['sheet_id']}"
             f"/values/{quote(range_ref)}", token)
    if r.status_code != 200:
        raise SheetError(f"Sheets API HTTP {r.status_code}: {r.text[:200]}")
    return _json(r, "Sheets API").get("values") or []


def _values_from_xlsx(sheet_cfg, token):
    r = _get(f"https://www.googleapis.com/drive/v3/files/{sheet_cfg['sheet_id']}"
             "?alt=media&supportsAllDrives=true", token, timeout=120)
    if r.status_code != 200:
        raise SheetError(f"Drive download HTTP {r.status_code}: {r.text[:200]}")
    try:
        import openpyxl
        from openpyxl.utils import column_index_from_string
        wb = openpyxl.load_workbook(BytesIO(r.content), read_only=True, data_only=True)
    except Exception as e:
        raise SheetError(f"could not parse the downloaded spreadsheet: {e}") from e
    tab = sheet_cfg["outreach_tab"]
    if tab not in wb.sheetnames:
        raise SheetError(
            f"tab '{tab}' not found in the outreach file — it has: {', '.join(wb.sheetnames)}")
    try:
        idx = column_index_from_string(sheet_cfg["handle_column"])
    except ValueError as e:
        raise SheetError(f"handle_column {sheet_cfg['handle_column']!r} is not a column letter") from e
    ws = wb[tab]
    return [[row[0]] for row in ws.iter_rows(min_col=idx, max_col=idx, values_only=True)]


def contacted_handles(sheet_cfg):
    sa_path = APP_DIR / sheet_cfg["service_account_file"]
    if not sa_path.exists():
        raise SheetError(
            f"service account file missing: {sa_path.name} — see README 'Sheet access' "
            "for how to create it and share the sheet with it.")
    try:
        from google.oauth2 import service_account
        from google.auth.transport.requests import Request as GoogleRequest
        creds = service_account.Credentials.from_service_account_file(str(sa_path), scopes=_SCOPES)
        creds.refresh(GoogleRequest())
    except Exception as e:
        raise SheetError(f"Google auth failed: {e}") from e

    mime = _file_mime(sheet_cfg["sheet_id"], creds.token)
    if mime == _GOOGLE_SHEET_MIME:
        values = _values_from_sheets_api(sheet_cfg, creds.token)
    elif mime == _XLSX_MIME:
        values = _values_from_xlsx(sheet_cfg, creds.token)
    else:
        raise SheetError(
            f"unsupported outreach file type {mime!r} — expected a native Google Sheet "
            "or an .xlsx file in Drive.")
    where = f"{sheet_cfg['outreach_tab']}!{sheet_cfg['handle_column']}"
    return parse_contacted(values, sheet_cfg["handle_header"], where)
